#!/usr/bin/env python3
"""
ebi_baseline_treelevel_yield_effect.py: the aggregate CV R^2 comparison
(compare_ebi_wholeimage_baseline.py) showed a small mean improvement from the
ensemble EBI, but that is a population-level summary. Per request: does swapping
the EBI baseline change yield PREDICTIONS AT THE INDIVIDUAL TREE LEVEL, i.e. would
two people using canonical vs whole-image EBI actually disagree about any specific
tree, not just about the average fit quality?

Reuses the exact same pooled 2022-2023 yield-model data pipeline as
compare_ebi_wholeimage_baseline.py (same n=167 tree-years, same base model:
cultivar + EBI x cultivar + climate + Growth_April). Fits the model in-sample
(full sample, not CV folds -- the question here is about the fitted relationship
and its predictions, not out-of-fold accuracy, which was already tested) under
each of the three EBI variants, then compares the per-tree fitted (predicted)
yield directly: correlation, mean/max absolute kg difference, rank stability
(would the "top" or "bottom" trees change), and which trees move the most.

-> Results_Analysis/08_UEF53_Rerun/EBI_Baseline_TreeLevel_Yield_Effect.json + figure.
"""
import os, sys, csv, math, json, sqlite3, struct
from collections import defaultdict
import numpy as np
import openpyxl
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from geo_guardrails import find_thesis_root, load_master

BASE = find_thesis_root()
MOS = os.path.join(BASE, "4band_mosaic")
GDIR = os.path.join(BASE, "Trees_Data_Survey", "Field_Data_Yield_GPKGs")
OUT = os.path.join(BASE, "Results_Analysis", "08_UEF53_Rerun")
WHOLEIMG_PATH = os.path.join(MOS, "Master_Trees_WholeImageBaseline_v3.xlsx")
PHYS_FIELDS = ["SWP_April", "Growth_April", "SWP_MayJune", "Growth_MayJune", "CNC_June", "SWP_June", "Growth_June"]
PHYS_GPKG_COLS = ["SWP_April", "Growth__April", "SWP_May.June", "Growth_May.June", "CNC_June", "SWP_June", "Growth_June"]
MATCH_THRESHOLD_M = 3.0

if not os.path.exists(WHOLEIMG_PATH):
    print(f"NOT FOUND: {WHOLEIMG_PATH}"); sys.exit(1)

M = load_master()
cultivar = {r["Tree_ID"]: r.get("cultivar") for r in M}
by_tid = {r["Tree_ID"]: r for r in M}

wb = openpyxl.load_workbook(WHOLEIMG_PATH, read_only=True)
ws = wb[wb.sheetnames[0]]
rows_ = list(ws.iter_rows(values_only=True))
header = rows_[0]; idx = {h: i for i, h in enumerate(header)}
wholeimg = {}
for r in rows_[1:]:
    tid = r[idx["Tree_ID"]]
    wholeimg[tid] = {h: r[idx[h]] for h in header if h not in ("Tree_ID", "Zone_Value")}


def zc(v):
    v = np.asarray(v, float); s = np.nanstd(v); return (v - np.nanmean(v)) / (s if s > 0 else 1)


def gpkg_geom_centroid(blob):
    flags = blob[3]; envelope_ind = (flags >> 1) & 0x07
    env_bytes = {0: 0, 1: 32, 2: 48, 3: 48, 4: 64}[envelope_ind]
    wkb = blob[8 + env_bytes:]; endian = "<" if wkb[0] == 1 else ">"
    pos = 1
    geom_type = struct.unpack_from(endian + "I", wkb, pos)[0]; pos += 4
    base_type = geom_type % 1000
    xs, ys = [], []
    if base_type == 6:
        npoly = struct.unpack_from(endian + "I", wkb, pos)[0]; pos += 4
        for _ in range(npoly):
            pos += 1; struct.unpack_from(endian + "I", wkb, pos)[0]; pos += 4
            nrings = struct.unpack_from(endian + "I", wkb, pos)[0]; pos += 4
            for r in range(nrings):
                npts = struct.unpack_from(endian + "I", wkb, pos)[0]; pos += 4
                for _p in range(npts):
                    x, y = struct.unpack_from(endian + "dd", wkb, pos); pos += 16
                    if r == 0: xs.append(x); ys.append(y)
    elif base_type == 3:
        nrings = struct.unpack_from(endian + "I", wkb, pos)[0]; pos += 4
        for r in range(nrings):
            npts = struct.unpack_from(endian + "I", wkb, pos)[0]; pos += 4
            for _p in range(npts):
                x, y = struct.unpack_from(endian + "dd", wkb, pos); pos += 16
                if r == 0: xs.append(x); ys.append(y)
    return sum(xs) / len(xs), sum(ys) / len(ys)


def load_phys(year):
    path = os.path.join(GDIR, f"Yield_with_clustering_{year}.gpkg")
    con = sqlite3.connect(path); cur = con.cursor()
    cur.execute("SELECT table_name FROM gpkg_geometry_columns")
    tbl = cur.fetchone()[0]
    cols = ",".join(f'"{c}"' for c in PHYS_GPKG_COLS)
    cur.execute(f'SELECT "Row","Plot",cultivar,{cols},geom FROM "{tbl}"')
    rows2 = cur.fetchall(); con.close()
    out = []
    for row_ in rows2:
        plot = row_[1]; vals = row_[3:3 + len(PHYS_FIELDS)]; blob = row_[-1]
        if plot != "A" or any(v is None for v in vals): continue
        x, y = gpkg_geom_centroid(blob)
        d = {"x": x, "y": y}
        for f, v in zip(PHYS_FIELDS, vals): d[f] = float(v)
        out.append(d)
    return out


PHYS = {y: load_phys(y) for y in [2022, 2023]}
mx = np.array([r.get("X_UTM") for r in M], float); my = np.array([r.get("Y_UTM") for r in M], float)
phys_match = {}
for y in [2022, 2023]:
    cand = []
    for i, r in enumerate(PHYS[y]):
        d = np.sqrt((mx - r["x"]) ** 2 + (my - r["y"]) ** 2)
        j = int(np.argmin(d))
        if d[j] <= MATCH_THRESHOLD_M: cand.append((d[j], i, j))
    cand.sort(); ui = set(); uj = set(); m = {}
    for d, i, j in cand:
        if i in ui or j in uj: continue
        m[j] = PHYS[y][i]; ui.add(i); uj.add(j)
    phys_match[y] = m

yc = list(csv.DictReader(open(os.path.join(BASE, "Trees_Data_Survey", "kedma_plot_a_yield.csv"), encoding="utf-8-sig")))
ymeas = defaultdict(dict); coord = {}
for r in yc:
    ymeas[r["tree_id"]][int(r["year"])] = float(r["net_kernel_yield_per_tree_kg"])
    coord[r["tree_id"]] = (float(r["latitude"]), float(r["longitude"]))
mlat = np.array([r.get("Latitude") for r in M], float); mlon = np.array([r.get("Longitude") for r in M], float)


def hav(a, b, d, e):
    p = math.pi / 180
    x = math.sin((d - a) * p / 2) ** 2 + math.cos(a * p) * math.cos(d * p) * math.sin((e - b) * p / 2) ** 2
    return 2 * 6371000 * math.asin(math.sqrt(max(0, x)))


cand = []
for t, (la, lo) in coord.items():
    ds = np.array([hav(la, lo, a, b) for a, b in zip(mlat, mlon)])
    for j in np.where(ds <= 8)[0]: cand.append((float(ds[j]), t, int(j)))
cand.sort(); uy = set(); um = set(); ymap = {}
for d, t, j in cand:
    if t in uy or j in um: continue
    ymap[t] = j; uy.add(t); um.add(j)

rowsP = [(t, j, yr) for t, j in ymap.items() for yr in (2022, 2023)
         if ymeas[t].get(yr) is not None and M[j].get(f"EBI_Norm_{yr}") is not None and j in phys_match[yr]
         and M[j]["Tree_ID"] in wholeimg and wholeimg[M[j]["Tree_ID"]].get(f"EBI_Norm_wholeimg_{yr}") is not None]
nP = len(rowsP)
print(f"pooled yield-model sample: n={nP}")

yP = np.array([ymeas[t][yr] for t, j, yr in rowsP])
uefP = np.array([1.0 if cultivar[M[j]["Tree_ID"]] == "UEF" else 0.0 for t, j, yr in rowsP])
ebiP_canon = zc([M[j].get(f"EBI_Norm_{yr}") for t, j, yr in rowsP])
ebiP_whole = zc([wholeimg[M[j]["Tree_ID"]].get(f"EBI_Norm_wholeimg_{yr}") for t, j, yr in rowsP])
ebiP_combo = (ebiP_canon + ebiP_whole) / 2.0
CP = {y: float(np.nanmean([r.get(f"Chill_Portions_{y}") for r in M if r.get(f"Chill_Portions_{y}") is not None])) for y in [2022, 2023]}
climP = zc([CP[yr] for t, j, yr in rowsP])
oneP = np.ones(nP)
phys_vals = {}
for f in PHYS_FIELDS:
    phys_vals[f] = zc([phys_match[yr][j][f] for t, j, yr in rowsP])
gaP = phys_vals["Growth_April"]

# ---------------- fit in-sample (full-sample OLS), one model per EBI variant ----------------
fitted = {}
for label, ebi_variant in [("canonical", ebiP_canon), ("wholeimg", ebiP_whole), ("combined", ebiP_combo)]:
    X = np.column_stack([oneP, uefP, ebi_variant, ebi_variant * uefP, climP, gaP])
    beta, *_ = np.linalg.lstsq(X, yP, rcond=None)
    fitted[label] = X @ beta

pred_canon = fitted["canonical"]; pred_whole = fitted["wholeimg"]; pred_combo = fitted["combined"]

# ---------------- per-tree-year comparison ----------------
diff_wc = pred_whole - pred_canon          # whole-image minus canonical, per row (kg)
diff_cc = pred_combo - pred_canon          # combined minus canonical, per row (kg)

def rank_corr(a, b):
    ra = np.argsort(np.argsort(a)).astype(float); rb = np.argsort(np.argsort(b)).astype(float)
    return float(np.corrcoef(ra, rb)[0, 1])


pearson_wc = float(np.corrcoef(pred_canon, pred_whole)[0, 1])
spearman_wc = rank_corr(pred_canon, pred_whole)
pearson_cc = float(np.corrcoef(pred_canon, pred_combo)[0, 1])
spearman_cc = rank_corr(pred_canon, pred_combo)

mae_wc = float(np.abs(diff_wc).mean()); max_wc = float(np.abs(diff_wc).max())
mae_cc = float(np.abs(diff_cc).mean()); max_cc = float(np.abs(diff_cc).max())
yield_mean = float(yP.mean()); yield_sd = float(yP.std())

print(f"\npred_canonical vs pred_wholeimg: Pearson r={pearson_wc:.4f} Spearman rho={spearman_wc:.4f} "
      f"MAE={mae_wc:.4f} kg  max|diff|={max_wc:.4f} kg")
print(f"pred_canonical vs pred_combined: Pearson r={pearson_cc:.4f} Spearman rho={spearman_cc:.4f} "
      f"MAE={mae_cc:.4f} kg  max|diff|={max_cc:.4f} kg")
print(f"(context: measured yield mean={yield_mean:.3f} kg, SD={yield_sd:.3f} kg)")

# does the TOP-10% / BOTTOM-10% predicted-yield tree set change?
k = max(1, int(round(0.1 * nP)))
top_canon = set(np.argsort(pred_canon)[-k:])
top_whole = set(np.argsort(pred_whole)[-k:])
top_combo = set(np.argsort(pred_combo)[-k:])
bot_canon = set(np.argsort(pred_canon)[:k])
bot_whole = set(np.argsort(pred_whole)[:k])
bot_combo = set(np.argsort(pred_combo)[:k])
top_overlap_wc = len(top_canon & top_whole) / k
top_overlap_cc = len(top_canon & top_combo) / k
bot_overlap_wc = len(bot_canon & bot_whole) / k
bot_overlap_cc = len(bot_canon & bot_combo) / k
print(f"\ntop-10% ({k} rows) overlap: canonical vs wholeimg={top_overlap_wc:.2f}, canonical vs combined={top_overlap_cc:.2f}")
print(f"bottom-10% overlap: canonical vs wholeimg={bot_overlap_wc:.2f}, canonical vs combined={bot_overlap_cc:.2f}")

# biggest movers (whole-image vs canonical), with cultivar breakdown
order = np.argsort(-np.abs(diff_wc))[:10]
biggest_movers = []
for i in order:
    t, j, yr = rowsP[i]
    biggest_movers.append({"tree": t, "year": yr, "cultivar": cultivar[M[j]["Tree_ID"]],
                            "pred_canonical_kg": round(float(pred_canon[i]), 3),
                            "pred_wholeimg_kg": round(float(pred_whole[i]), 3),
                            "diff_kg": round(float(diff_wc[i]), 3),
                            "measured_kg": round(float(yP[i]), 3)})

# is the divergence bigger for one cultivar than the other?
uef_mask = uefP == 1.0
mae_wc_uef = float(np.abs(diff_wc[uef_mask]).mean()); mae_wc_53 = float(np.abs(diff_wc[~uef_mask]).mean())

RESULTS = {
    "note": "Per-tree-year (n=167) comparison of FITTED (in-sample) predicted yield under the "
            "canonical, whole-image, and combined EBI variants, same pooled 2022-2023 model as "
            "compare_ebi_wholeimage_baseline.py. Answers whether the choice of baseline reshuffles "
            "which specific trees look good/bad, not just the aggregate CV R^2.",
    "n": nP,
    "measured_yield_context": {"mean_kg": round(yield_mean, 3), "sd_kg": round(yield_sd, 3)},
    "canonical_vs_wholeimg": {
        "pearson_r_of_fitted_predictions": round(pearson_wc, 4),
        "spearman_rho_of_fitted_predictions": round(spearman_wc, 4),
        "mean_abs_diff_kg": round(mae_wc, 4), "max_abs_diff_kg": round(max_wc, 4),
        "mean_abs_diff_kg_UEF": round(mae_wc_uef, 4), "mean_abs_diff_kg_cultivar53": round(mae_wc_53, 4),
        "top10pct_predicted_yield_tree_overlap": round(top_overlap_wc, 3),
        "bottom10pct_predicted_yield_tree_overlap": round(bot_overlap_wc, 3),
    },
    "canonical_vs_combined": {
        "pearson_r_of_fitted_predictions": round(pearson_cc, 4),
        "spearman_rho_of_fitted_predictions": round(spearman_cc, 4),
        "mean_abs_diff_kg": round(mae_cc, 4), "max_abs_diff_kg": round(max_cc, 4),
        "top10pct_predicted_yield_tree_overlap": round(top_overlap_cc, 3),
        "bottom10pct_predicted_yield_tree_overlap": round(bot_overlap_cc, 3),
    },
    "biggest_movers_canonical_vs_wholeimg": biggest_movers,
}
with open(os.path.join(OUT, "EBI_Baseline_TreeLevel_Yield_Effect.json"), "w") as f:
    json.dump(RESULTS, f, indent=2, default=str)
print("\nsaved json")

# ---------------- figure ----------------
fig, axes = plt.subplots(1, 2, figsize=(12, 4.8))
ax = axes[0]
ax.scatter(pred_canon, pred_whole, s=18, alpha=0.6, color="#16A085", label=f"wholeimg (r={pearson_wc:.3f})")
ax.scatter(pred_canon, pred_combo, s=18, alpha=0.4, color="#8E44AD", label=f"combined (r={pearson_cc:.3f})")
lims = [min(pred_canon.min(), pred_whole.min(), pred_combo.min()), max(pred_canon.max(), pred_whole.max(), pred_combo.max())]
ax.plot(lims, lims, "k--", lw=1)
ax.set_xlabel("predicted yield, canonical EBI (kg)"); ax.set_ylabel("predicted yield, alternative variant (kg)")
ax.set_title("Per-tree-year predicted yield:\ncanonical vs alternative baselines", fontsize=10)
ax.legend(fontsize=8)

ax = axes[1]
ax.hist(diff_wc, bins=25, color="#16A085", alpha=0.6, label=f"wholeimg-canonical (MAE={mae_wc:.3f} kg)")
ax.hist(diff_cc, bins=25, color="#8E44AD", alpha=0.5, label=f"combined-canonical (MAE={mae_cc:.3f} kg)")
ax.axvline(0, color="k", ls="--", lw=1)
ax.set_xlabel("predicted-yield difference vs canonical (kg)"); ax.set_ylabel("count (of 167 tree-years)")
ax.set_title("Distribution of per-tree-year\nprediction shift", fontsize=10)
ax.legend(fontsize=8)
plt.tight_layout()
plt.savefig(os.path.join(OUT, "EBI_Baseline_TreeLevel_Yield_Effect.png"), dpi=130)
print("saved figure")
