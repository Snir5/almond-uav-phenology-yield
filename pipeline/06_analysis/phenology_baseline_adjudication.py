#!/usr/bin/env python3
"""
phenology_baseline_adjudication.py: the recommended follow-up from section 13.6 --
adjudicate whether canonical (canopy-only) or whole-image-baseline EBI better
matches INDEPENDENT ground truth: the field phenology surveys' "% bloom open on the
UAV flight date" per cultivar-year (Kadma_Phenology_by_CultivarDate_2022_2024.csv +
the 2021 20-tree bloom survey). This ground truth does not depend on the mosaic or
its radiometric correction at all, so it is a fair referee between the two EBI
variants.

Reuses phenology_flight_timing.py's own field-curve loading and flight-date
interpolation exactly (same FLIGHT dict, same curve-building logic), so the ground
truth values are identical to the already-published Figure 22/§16 analysis; only the
EBI side of the comparison is extended to the whole-image and combined variants.

-> Results_Analysis/08_UEF53_Rerun/Phenology_Baseline_Adjudication.json + figure.
"""
import os, sys, csv, json
import numpy as np
from datetime import datetime
import openpyxl
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from geo_guardrails import find_thesis_root, load_master
from stats_utils import r_pvalue

BASE = find_thesis_root()
MOS = os.path.join(BASE, "4band_mosaic")
FIG = os.path.join(BASE, "Results_Analysis", "08_UEF53_Rerun")
WHOLEIMG_PATH = os.path.join(MOS, "Master_Trees_WholeImageBaseline_v3.xlsx")
FLIGHT = {2021: 66, 2022: 61, 2023: 60, 2024: 60}   # identical to phenology_flight_timing.py

if not os.path.exists(WHOLEIMG_PATH):
    print(f"NOT FOUND: {WHOLEIMG_PATH}")
    print("Run apply_fixed_zones_yearly_v3_wholeimage_baseline.py first.")
    sys.exit(1)

M = load_master()
cult = np.array([r.get("cultivar") for r in M])


def col(c):
    return np.array([r.get(c) if r.get(c) is not None else np.nan for r in M], float)


ebimean_canon = {y: {c: float(np.nanmean(col(f"EBI_Norm_{y}")[cult == c])) for c in ["53", "UEF"]}
                  for y in [2021, 2022, 2023, 2024]}

# ---------------- load the whole-image-baseline workbook, same tree set as canonical ----------------
wb = openpyxl.load_workbook(WHOLEIMG_PATH, read_only=True)
ws = wb[wb.sheetnames[0]]
rows_ = list(ws.iter_rows(values_only=True))
header = rows_[0]; idx = {h: i for i, h in enumerate(header)}
wholeimg = {}
for r in rows_[1:]:
    tid = r[idx["Tree_ID"]]
    wholeimg[tid] = {h: r[idx[h]] for h in header if h not in ("Tree_ID", "Zone_Value")}

cult_by_tid = {r["Tree_ID"]: r.get("cultivar") for r in M}
ebimean_whole = {}
for y in [2021, 2022, 2023, 2024]:
    ebimean_whole[y] = {}
    for c in ["53", "UEF"]:
        vals = [wholeimg[tid].get(f"EBI_Norm_wholeimg_{y}") for tid in wholeimg
                if cult_by_tid.get(tid) == c and wholeimg[tid].get(f"EBI_Norm_wholeimg_{y}") is not None]
        ebimean_whole[y][c] = float(np.mean(vals)) if vals else None

# combined = simple average of the two means (both already on the same 0-1 EBI_Norm scale,
# comparable because both are min-max normalized to their OWN reference-year (2024) 1st/99th
# percentile range -- an approximation stated plainly, not a pooled-sample z-score as in Part C
# of the yield-model comparison, since here we only have year x cultivar CELL MEANS, not trees)
ebimean_combo = {y: {c: (ebimean_canon[y][c] + ebimean_whole[y][c]) / 2.0
                     if ebimean_whole[y].get(c) is not None else None
                     for c in ["53", "UEF"]} for y in [2021, 2022, 2023, 2024]}


# ---------------- identical field-curve loading to phenology_flight_timing.py ----------------
def doy(s):
    return datetime.strptime(s, "%d/%m/%Y").timetuple().tm_yday


def norm_c(c):
    return "UEF" if ("UEF" in c or "Um" in c) else c


curve = {}
for r in csv.DictReader(open(os.path.join(BASE, "Trees_Data_Survey", "Phenology_Surveys",
                                           "Kadma_Phenology_by_CultivarDate_2022_2024.csv"), encoding="utf-8-sig")):
    c = norm_c(r["Cultivar"])
    if c in ("53", "UEF"):
        curve.setdefault((int(r["Year"]), c), []).append((doy(r["Date"]), float(r["Mean_Percent_Open"])))
p21 = {}
for r in csv.DictReader(open(os.path.join(BASE, "Trees_Data_Survey", "Bloom_Survey_2021",
                                           "Phenology_Survey_2021_20trees.csv"), encoding="utf-8-sig")):
    c = "53" if "53" in r["Cultivar"] else "UEF"
    p21.setdefault((c, doy(r["Date"])), []).append(float(r["Flowers"]))
for (c, d), v in p21.items():
    curve.setdefault((2021, c), []).append((d, float(np.mean(v))))


def interp(pts, x):
    pts = sorted(pts); xs = [p[0] for p in pts]; ys = [p[1] for p in pts]
    return float(np.interp(x, xs, ys))


rows = []
for y in [2021, 2022, 2023, 2024]:
    for c in ["53", "UEF"]:
        if (y, c) not in curve:
            continue
        pts = sorted(curve[(y, c)])
        po = interp(pts, FLIGHT[y])
        rows.append({"year": y, "cultivar": c, "pct_open_at_flight": round(po, 1),
                     "EBI_canonical": round(ebimean_canon[y][c], 4),
                     "EBI_wholeimg": round(ebimean_whole[y][c], 4) if ebimean_whole[y].get(c) is not None else None,
                     "EBI_combined": round(ebimean_combo[y][c], 4) if ebimean_combo[y].get(c) is not None else None})

po = np.array([r["pct_open_at_flight"] for r in rows])
eb_c = np.array([r["EBI_canonical"] for r in rows])
eb_w = np.array([r["EBI_wholeimg"] for r in rows])
eb_x = np.array([r["EBI_combined"] for r in rows])
n = len(rows)


def corr_report(a, b):
    r = float(np.corrcoef(a, b)[0, 1])
    return {"r": round(r, 3), "p": round(float(r_pvalue(r, len(a))), 4)}


rep_canon = corr_report(po, eb_c)
rep_whole = corr_report(po, eb_w)
rep_combo = corr_report(po, eb_x)

print(f"n={n} (year x cultivar cells with both field survey and EBI data)")
print(f"EBI_canonical vs field pct_open_at_flight: r={rep_canon['r']} p={rep_canon['p']}")
print(f"EBI_wholeimg  vs field pct_open_at_flight: r={rep_whole['r']} p={rep_whole['p']}")
print(f"EBI_combined  vs field pct_open_at_flight: r={rep_combo['r']} p={rep_combo['p']}")

RESULTS = {
    "note": "Adjudicates canonical vs whole-image-baseline EBI against INDEPENDENT ground truth "
            "(field-measured % bloom open at the UAV flight date, per cultivar-year), which does not "
            "depend on the mosaic or its radiometric correction. Reuses phenology_flight_timing.py's "
            "own field-curve loading and flight-date interpolation exactly, so ground-truth values "
            "match the already-published §16 analysis; only the EBI side is extended.",
    "n": n,
    "caveat": "n=7 (year x cultivar cells; cultivar 53 has no 2024 field survey). Correlations at this "
              "n have wide confidence intervals; read directionally, not as a settled verdict.",
    "rows": rows,
    "EBI_canonical_vs_field_r": rep_canon,
    "EBI_wholeimg_vs_field_r": rep_whole,
    "EBI_combined_vs_field_r": rep_combo,
}
with open(os.path.join(FIG, "Phenology_Baseline_Adjudication.json"), "w") as f:
    json.dump(RESULTS, f, indent=2, default=str)
print("\nsaved json")

# ---------------- figure ----------------
fig, axes = plt.subplots(1, 3, figsize=(15, 4.8), sharex=True, sharey=True)
for ax, (label, eb, rep) in zip(axes, [("Canonical (canopy-only)", eb_c, rep_canon),
                                        ("Whole-image baseline", eb_w, rep_whole),
                                        ("Combined (ensemble)", eb_x, rep_combo)]):
    ax.scatter(po, eb, color="#C0392B", s=60, zorder=3)
    for r in rows:
        ax.annotate(f"{r['year']}-{r['cultivar']}", (r["pct_open_at_flight"], r[f"EBI_{'canonical' if 'Canonical' in label else ('wholeimg' if 'Whole' in label else 'combined')}"]),
                    fontsize=7, xytext=(4, 4), textcoords="offset points")
    if len(set(po)) > 1:
        m, b_ = np.polyfit(po, eb, 1)
        xs = np.linspace(po.min(), po.max(), 20)
        ax.plot(xs, m * xs + b_, "--", color="#16A085", lw=1.5)
    ax.set_title(f"{label}\nr={rep['r']}, p={rep['p']}", fontsize=10)
    ax.set_xlabel("field % bloom open at flight date")
axes[0].set_ylabel("mean EBI_Norm (year x cultivar)")
fig.suptitle("Which EBI baseline better tracks independent field bloom phenology?", fontweight="bold", fontsize=12)
plt.tight_layout()
plt.savefig(os.path.join(FIG, "Phenology_Baseline_Adjudication.png"), dpi=130)
print("saved figure")
