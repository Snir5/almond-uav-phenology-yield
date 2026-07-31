#!/usr/bin/env python3
"""
build_master_with_canopy_features.py: per request, adds the canopy/bloom feature
bank (section 38) to a NEW master file (does not overwrite the canonical
Master_Trees_Extended.xlsx, per project convention). Copies every column from the
canonical master, then appends, for each of 2021-2024:
  CanopyArea_m2_{y}, CanopyCoverFraction_{y}, ShadowFraction_{y} (from
    Master_Trees_CanopyStructure_v5.xlsx)
  BloomFraction_{y}, BrightFraction_{y} (from Master_Trees_BloomFraction_v3.xlsx,
    the Otsu pixel-threshold-crossing-percentage features)
  EBI_density_{y}, BloomVolume_{y}, BloomPixelVolume_{y}, EBI_x_BloomFraction_{y},
    EBI_x_CanopyCover_{y}, BrightFraction_x_CanopyArea_{y} (engineered, section 38)
Matched by Tree_ID (both source files use the same Tree_ID scheme as the master;
their own X_UTM/Y_UTM columns are on a different footing, confirmed in section 38).
-> 4band_mosaic/Master_Trees_Extended_CanopyFeatures.xlsx
"""
import os, sys, json
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from geo_guardrails import find_thesis_root

BASE = find_thesis_root()
MDIR = os.path.join(BASE, "4band_mosaic")


def load_xlsx_by_id(path):
    wb = openpyxl.load_workbook(path, read_only=True); ws = wb.active
    rows = list(ws.iter_rows(values_only=True)); header = rows[0]; hidx = {h: i for i, h in enumerate(header)}
    return {r[hidx["Tree_ID"]]: {h: r[hidx[h]] for h in header} for r in rows[1:]}


canopy_by_id = load_xlsx_by_id(os.path.join(MDIR, "Master_Trees_CanopyStructure_v5.xlsx"))
bloomfrac_by_id = load_xlsx_by_id(os.path.join(MDIR, "Master_Trees_BloomFraction_v3.xlsx"))

wb_master = openpyxl.load_workbook(os.path.join(MDIR, "Master_Trees_Extended.xlsx"))
ws_master = wb_master.active
header = [c.value for c in next(ws_master.iter_rows(min_row=1, max_row=1))]
tree_id_col = header.index("Tree_ID") + 1

YEARS = [2021, 2022, 2023, 2024]
new_cols = []
for y in YEARS:
    new_cols += [f"CanopyArea_m2_{y}", f"CanopyCoverFraction_{y}", f"ShadowFraction_{y}",
                 f"BloomFraction_{y}", f"BrightFraction_{y}",
                 f"EBI_density_{y}", f"BloomVolume_{y}", f"BloomPixelVolume_{y}",
                 f"EBI_x_BloomFraction_{y}", f"EBI_x_CanopyCover_{y}", f"BrightFraction_x_CanopyArea_{y}"]

start_col = len(header) + 1
for i, name in enumerate(new_cols):
    ws_master.cell(row=1, column=start_col + i, value=name)

n_rows_written = 0
n_ebi_missing = {y: 0 for y in YEARS}
for row_i in range(2, ws_master.max_row + 1):
    tid = ws_master.cell(row=row_i, column=tree_id_col).value
    c = canopy_by_id.get(tid); b = bloomfrac_by_id.get(tid)
    for y in YEARS:
        ebi = ws_master.cell(row=row_i, column=header.index(f"EBI_Norm_{y}") + 1).value if f"EBI_Norm_{y}" in header else None
        col_offset = YEARS.index(y) * 11
        vals = [None] * 11
        if c is not None and b is not None:
            area = c.get(f"CanopyArea_m2_{y}"); cover = c.get(f"CanopyCoverFraction_{y}")
            shadow = c.get(f"ShadowFraction_{y}"); pix = c.get(f"CanopyPixelCount_{y}")
            bf = b.get(f"BloomFraction_{y}"); brf = b.get(f"BrightFraction_{y}")
            vals[0] = area; vals[1] = cover; vals[2] = shadow; vals[3] = bf; vals[4] = brf
            if all(v is not None for v in [area, bf, pix]) and ebi is not None:
                area_f = float(area); ebi_f = float(ebi); bf_f = float(bf); pix_f = float(pix)
                vals[5] = ebi_f / area_f if area_f > 0 else None
                vals[6] = ebi_f * area_f
                vals[7] = bf_f * pix_f
                vals[8] = ebi_f * bf_f
                vals[9] = (ebi_f * float(cover)) if cover is not None else None
                vals[10] = (float(brf) * area_f) if brf is not None else None
            else:
                if ebi is None: n_ebi_missing[y] += 1
        for k, v in enumerate(vals):
            ws_master.cell(row=row_i, column=start_col + col_offset + k, value=v)
    n_rows_written += 1

out_path = os.path.join(MDIR, "Master_Trees_Extended_CanopyFeatures.xlsx")
wb_master.save(out_path)
print(f"wrote {n_rows_written} rows, {len(new_cols)} new columns, to {out_path}")
print(f"EBI missing (so engineered features left blank) by year: {n_ebi_missing}")
print(f"trees matched to canopy structure file: {sum(1 for tid in [ws_master.cell(row=r, column=tree_id_col).value for r in range(2, ws_master.max_row+1)] if canopy_by_id.get(tid) is not None)}")
